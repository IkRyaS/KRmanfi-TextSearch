import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from krmanfi_search.infrastructure.search import extractors
from krmanfi_search.infrastructure.search.extractors import (
    ExtractorRegistry,
    PdfExtractor,
    decode_text,
    extract_docx,
    extract_plain,
    extract_pptx,
    extract_xlsx,
    image_coverage,
    page_needs_ocr,
    page_zoom,
    pdf_page_count,
    render_page_image,
)
from krmanfi_search.shared.exceptions import OperationCancelledError
from krmanfi_search.shared.types import CancelToken
from tests.helpers import make_image_pdf, make_text_pdf, solid_png


def cancelled():
    token = CancelToken()
    token.cancel()
    return token


def test_decode_utf8_with_bom():
    assert decode_text("﻿привет".encode()) == "привет"


def test_decode_utf16_both_byte_orders():
    assert decode_text("привет".encode("utf-16")) == "привет"
    assert decode_text(b"\xfe\xff" + "ok".encode("utf-16-be")) == "ok"


def test_decode_cp1251_when_not_utf8():
    assert decode_text("Квартальный отчёт".encode("cp1251")) == "Квартальный отчёт"


def test_decode_falls_back_to_latin1():
    assert decode_text(b"\x98abc") == "\x98abc"


def test_plain_text_is_split_by_500_lines(tmp_path):
    path = tmp_path / "big.txt"
    path.write_text("\n".join(f"строка {n}" for n in range(1, 1201)), encoding="utf-8")
    chunks = extract_plain(path, CancelToken())
    assert [c.location for c in chunks] == ["строки 1–500", "строки 501–1000", "строки 1001–1200"]


def test_blank_text_file_gives_no_chunks(tmp_path):
    path = tmp_path / "blank.txt"
    path.write_text("\n   \n\n", encoding="utf-8")
    assert extract_plain(path, CancelToken()) == []


def test_empty_file_gives_no_chunks(tmp_path):
    path = tmp_path / "empty.md"
    path.write_bytes(b"")
    assert extract_plain(path, CancelToken()) == []


def test_docx_paragraphs_tables_and_footers(tmp_path):
    document = Document()
    document.add_paragraph("Первый абзац")
    document.add_paragraph("   ")
    document.add_paragraph("Второй абзац")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "ячейка А"
    table.rows[0].cells[1].text = "ячейка Б"
    document.sections[0].footer.paragraphs[0].text = "нижний колонтитул"
    document.sections[0].header.paragraphs[0].text = "верхний колонтитул"
    path = tmp_path / "a.docx"
    document.save(path)
    chunks = extract_docx(path, CancelToken())
    assert chunks[0].location == "абзацы 1–3"
    assert "ячейка А\tячейка Б" in chunks[0].text
    assert "   " not in chunks[0].text.splitlines()
    assert chunks[-1].location == "колонтитулы"
    assert "нижний колонтитул" in chunks[-1].text and "верхний колонтитул" in chunks[-1].text


def test_docx_without_footers_has_no_footer_chunk(tmp_path):
    document = Document()
    document.add_paragraph("только текст")
    path = tmp_path / "plain.docx"
    document.save(path)
    assert [c.location for c in extract_docx(path, CancelToken())] == ["абзацы 1–1"]


def test_docx_splits_many_paragraphs(tmp_path):
    document = Document()
    for number in range(120):
        document.add_paragraph(f"абзац {number}")
    path = tmp_path / "long.docx"
    document.save(path)
    labels = [c.location for c in extract_docx(path, CancelToken())]
    assert labels == ["абзацы 1–50", "абзацы 51–100", "абзацы 101–120"]


def test_xlsx_chunk_per_sheet_skips_empty_cells_and_sheets(tmp_path):
    workbook = Workbook()
    first = workbook.active
    first.title = "Платежи"
    first.append(["договор", None, 42])
    first.append([None, None, None])
    workbook.create_sheet("Пусто")
    third = workbook.create_sheet("Итоги")
    third.append(["всего", 100])
    path = tmp_path / "a.xlsx"
    workbook.save(path)
    chunks = extract_xlsx(path, CancelToken())
    assert [(c.location, c.page) for c in chunks] == [("лист «Платежи»", 1), ("лист «Итоги»", 3)]
    assert chunks[0].text == "договор\t42"


def test_xlsx_cancel_is_honoured(tmp_path):
    workbook = Workbook()
    workbook.active.append(["a"])
    path = tmp_path / "a.xlsx"
    workbook.save(path)
    with pytest.raises(OperationCancelledError):
        extract_xlsx(path, cancelled())


def test_pptx_text_notes_and_tables(tmp_path):
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[5])
    first.shapes.title.text = "Заголовок слайда"
    first.notes_slide.notes_text_frame.text = "заметка докладчика"
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = second.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "левая"
    table.cell(0, 1).text = "правая"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "a.pptx"
    presentation.save(path)
    chunks = extract_pptx(path, CancelToken())
    assert [(c.location, c.page) for c in chunks] == [("слайд 1", 1), ("слайд 2", 2)]
    assert "заметка докладчика" in chunks[0].text and "Заголовок слайда" in chunks[0].text
    assert "левая" in chunks[1].text and "правая" in chunks[1].text


def test_pptx_cancel_is_honoured(tmp_path):
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "a.pptx"
    presentation.save(path)
    with pytest.raises(OperationCancelledError):
        extract_pptx(path, cancelled())


def test_pdf_text_layer_pages(tmp_path):
    path = make_text_pdf(tmp_path / "t.pdf", ["first page", "", "third page"])
    chunks = PdfExtractor().extract(path, CancelToken())
    assert [(c.location, c.page) for c in chunks] == [("стр. 1", 1), ("стр. 3", 3)]


@pytest.mark.parametrize(
    ("start", "stop", "expected"),
    [(0, 2, [1, 2]), (2, 4, [3, 4]), (3, None, [4, 5]), (4, 99, [5]), (7, None, []), (-5, 1, [1])],
)
def test_pdf_page_ranges(tmp_path, start, stop, expected):
    path = make_text_pdf(tmp_path / "t.pdf", [f"page {n}" for n in range(1, 6)])
    chunks = PdfExtractor().extract_pages(path, start, stop, CancelToken())
    assert [c.page for c in chunks] == expected


def test_encrypted_pdf_raises_readable_error(tmp_path):
    path = make_text_pdf(tmp_path / "e.pdf", ["secret"], password="pw")
    with pytest.raises(ValueError, match="зашифрован"):
        PdfExtractor().extract(path, CancelToken())


def test_pdf_extraction_stops_when_cancelled(tmp_path):
    path = make_text_pdf(tmp_path / "t.pdf", ["a", "b"])
    with pytest.raises(OperationCancelledError):
        PdfExtractor().extract(path, cancelled())


def test_pdf_page_count_variants(tmp_path):
    assert pdf_page_count(make_text_pdf(tmp_path / "t.pdf", ["a", "b", "c"])) == 3
    assert pdf_page_count(make_text_pdf(tmp_path / "e.pdf", ["a"], password="pw")) == 0
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really")
    assert pdf_page_count(broken) == 0
    assert pdf_page_count(tmp_path / "missing.pdf") == 0


def test_image_coverage_full_page_and_none(tmp_path):
    with pymupdf.open(str(make_image_pdf(tmp_path / "i.pdf"))) as document:
        assert image_coverage(document[0]) == pytest.approx(1.0, abs=0.01)
    with pymupdf.open(str(make_text_pdf(tmp_path / "t.pdf", ["x"]))) as document:
        assert image_coverage(document[0]) == 0.0


def test_page_needs_ocr_for_image_only_page(tmp_path):
    with pymupdf.open(str(make_image_pdf(tmp_path / "i.pdf"))) as document:
        assert page_needs_ocr(document[0], "")
        assert page_needs_ocr(document[0], "a" * 49)
        assert page_needs_ocr(document[0], "a" * 100)
        assert not page_needs_ocr(document[0], "a" * 300)


def test_page_needs_ocr_false_without_images(tmp_path):
    with pymupdf.open(str(make_text_pdf(tmp_path / "t.pdf", ["x"]))) as document:
        assert not page_needs_ocr(document[0], "")


def test_small_picture_with_real_text_needs_no_ocr(tmp_path):
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(10, 10, 60, 60), stream=solid_png(20, 20))
    assert not page_needs_ocr(page, "x" * 80)
    assert page_needs_ocr(page, "x" * 10)
    document.close()


def test_page_zoom_follows_dpi_and_is_limited(tmp_path):
    with pymupdf.open(str(make_text_pdf(tmp_path / "t.pdf", ["x"]))) as document:
        assert page_zoom(document[0], 144) == pytest.approx(2.0)
    big = pymupdf.open()
    page = big.new_page(width=5000, height=5000)
    assert page_zoom(page, 600) == pytest.approx(extractors.MAX_OCR_RENDER_PIXELS / 5000)
    big.close()


def test_rendered_page_is_a_pgm(tmp_path):
    with pymupdf.open(str(make_text_pdf(tmp_path / "t.pdf", ["x"]))) as document:
        assert render_page_image(document[0], 72).startswith(b"P5")


def test_ocr_text_is_added_to_image_page_and_marked(tmp_path):
    path = make_image_pdf(tmp_path / "i.pdf", pages=2)
    seen = []

    def fake_ocr(image):
        seen.append(image[:2])
        return "распознанный текст"

    chunks = PdfExtractor(fake_ocr).extract(path, CancelToken())
    assert [c.location for c in chunks] == ["стр. 1 (OCR)", "стр. 2 (OCR)"]
    assert seen == [b"P5", b"P5"]


def test_ocr_result_is_appended_to_existing_text(tmp_path):
    path = make_image_pdf(tmp_path / "i.pdf", text="short caption")
    chunks = PdfExtractor(lambda image: "из картинки").extract(path, CancelToken())
    assert "short caption" in chunks[0].text and "из картинки" in chunks[0].text


def test_blank_ocr_result_leaves_page_without_text(tmp_path):
    path = make_image_pdf(tmp_path / "i.pdf")
    assert PdfExtractor(lambda image: "  \n ").extract(path, CancelToken()) == []


def test_ocr_is_not_called_without_images(tmp_path):
    path = make_text_pdf(tmp_path / "t.pdf", ["plain text page"])

    def forbidden(image):
        raise AssertionError("OCR must not run")

    assert PdfExtractor(forbidden).extract(path, CancelToken())[0].location == "стр. 1"


def test_registry_dispatch_by_extension(tmp_path):
    registry = ExtractorRegistry()
    text = tmp_path / "a.log"
    text.write_text("строка", encoding="utf-8")
    assert registry.extract(text, "log", CancelToken(), None)[0].text == "строка"
    assert registry.extract(text, "exe", CancelToken(), None) == []
    pdf = make_text_pdf(tmp_path / "t.pdf", ["one", "two", "three"])
    assert len(registry.extract(pdf, "pdf", CancelToken(), None)) == 3
    assert [c.page for c in registry.extract(pdf, "pdf", CancelToken(), (1, 2))] == [2]
