import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from tests.conftest import PHRASE


def where(hits):
    return {
        (h.display_path.rsplit("!/", 1)[-1].split("/")[-1].split("\\")[-1], h.location)
        for h in hits
    }


@pytest.fixture
def mixed(tmp_path: Path):
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "a.txt").write_text(
        "Первый Квартальный Отчёт\nвторой квартальный отчет\n", encoding="utf-8"
    )
    (folder / "b.csv").write_text("id;имя\n1;квартальный отчёт\n", encoding="cp1251")
    doc = Document()
    doc.add_paragraph("абзац без совпадений")
    doc.add_paragraph("тут квартальный отчёт")
    doc.save(folder / "c.docx")
    book = Workbook()
    book.active.title = "Лист1"
    book.active.append(["квартальный отчёт"])
    book.save(folder / "d.xlsx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "заголовок"
    slide.placeholders[1].text = "квартальный отчёт"
    deck.save(folder / "e.pptx")
    inner = folder / "inner.zip"
    with zipfile.ZipFile(inner, "w") as zf:
        zf.writestr("x/f.txt", "квартальный отчёт во вложенном")
    with zipfile.ZipFile(folder / "outer.zip", "w") as zf:
        zf.write(inner, "in/inner.zip")
        zf.writestr("g.md", "квартальный отчёт")
    inner.unlink()
    return folder


def test_text_found_in_every_supported_format(run_search, mixed):
    hits, stats = run_search(mixed, ocr_scanned_pdfs=False)
    names = {name for name, _ in where(hits)}
    assert names == {"a.txt", "b.csv", "c.docx", "d.xlsx", "e.pptx", "f.txt", "g.md"}
    assert stats.failed == 0


def test_locations_are_labelled(run_search, mixed):
    hits, _ = run_search(mixed, ocr_scanned_pdfs=False)
    locations = {loc for _, loc in where(hits)}
    assert {"строки 1–2", "абзацы 1–2", "лист «Лист1»", "слайд 1"} <= locations


def test_every_occurrence_is_its_own_hit(run_search, mixed):
    hits, _ = run_search(mixed, ocr_scanned_pdfs=False)
    own = [h for h in hits if h.display_path.endswith("a.txt")]
    assert [h.occurrence for h in own] == [1, 2]
    assert {h.match_count for h in own} == {2}
    assert own[0].fragment == "Квартальный Отчёт"


def test_context_has_ellipsis_and_collapsed_spaces(run_search, tmp_path):
    (tmp_path / "t.txt").write_text(
        "слово " * 30 + "иголка   рядом " + "слово " * 30, encoding="utf-8"
    )
    hits, _ = run_search(tmp_path, "иголка", context_chars=10, ocr_scanned_pdfs=False)
    assert hits[0].context.startswith("…") and hits[0].context.endswith("…")
    assert "  " not in hits[0].context


def test_nested_archive_member_keeps_full_path(run_search, mixed):
    hits, _ = run_search(mixed, ocr_scanned_pdfs=False)
    deep = next(h for h in hits if h.display_path.endswith("f.txt"))
    assert deep.display_path.count("!/") == 2
    assert deep.real_path.name == "outer.zip"


def test_archives_can_be_switched_off(run_search, mixed):
    hits, _ = run_search(mixed, include_archives=False, ocr_scanned_pdfs=False)
    assert not any("!/" in h.display_path for h in hits)


def test_regex_mode(run_search, mixed):
    hits, _ = run_search(mixed, r"квартальн\w+ отч\w+", mode="regex", ocr_scanned_pdfs=False)
    assert len(hits) >= 7


def test_single_archive_as_root(run_search, mixed):
    hits, stats = run_search(mixed / "outer.zip", ocr_scanned_pdfs=False)
    assert {h.display_path.split("!/", 1)[1] for h in hits} == {"g.md", "in/inner.zip!/x/f.txt"}
    assert stats.archives == 1


def test_nothing_found_gives_clear_stats(run_search, mixed):
    hits, stats = run_search(mixed, "отсутствующая фраза", ocr_scanned_pdfs=False)
    assert hits == []
    assert stats.files_total > 0 and stats.hits == 0


def test_phrase_constant_is_cyrillic():
    assert PHRASE == "квартальный отчёт"
