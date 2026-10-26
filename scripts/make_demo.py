"""Создаёт демо-набор ``demo/``: синтетические файлы без личных данных.

Запуск: ``python scripts/make_demo.py``. Нужны пакеты из extra ``dev`` и шрифт с кириллицей
(DejaVu Sans, Liberation Sans или Arial; путь можно задать в ``KRMANFI_DEMO_FONT``).
"""

from __future__ import annotations

import os
import random
import shutil
import sys
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "demo"
PHRASE = "квартальный отчёт"
FONT_CANDIDATES = (
    os.environ.get("KRMANFI_DEMO_FONT", ""),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/Library/Fonts/Arial.ttf",
)
PAGE_W, PAGE_H = 1654, 2339


def find_font():
    """Путь к шрифту с кириллицей."""
    for candidate in FONT_CANDIDATES:
        if candidate and Path(candidate).is_file():
            return candidate
    sys.exit("Не найден шрифт с кириллицей: задайте KRMANFI_DEMO_FONT")


FONT = find_font()


def text_pdf(path: Path, pages: list[str]):
    """PDF с текстовым слоем: по строке-тексту на страницу."""
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((60, 100), text, fontname="dv", fontfile=FONT, fontsize=13)
    doc.save(str(path))
    doc.close()


def page_image(lines: list[str], photo: bool = False):
    """Страница-картинка с текстом; ``photo`` добавляет тень, шум и поворот."""
    image = Image.new("L", (PAGE_W, PAGE_H), 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, 54)
    y = 220
    for line in lines:
        draw.text((140, y), line, fill=20, font=font)
        y += 110
    if photo:
        rng = random.Random(7)
        pixels = image.load()
        assert pixels is not None
        for x in range(0, PAGE_W, 2):
            shade = int(60 * x / PAGE_W)
            for yy in range(0, PAGE_H, 2):
                value = pixels[x, yy]
                if isinstance(value, int):
                    pixels[x, yy] = max(0, value - shade - rng.randint(0, 10))
        image = image.filter(ImageFilter.GaussianBlur(0.8)).rotate(1.5, fillcolor=200)
    return image


def scan_pdf(path: Path, pages: list[list[str]], photo: bool = False, rotate: int = 0):
    """PDF из одних картинок (без текстового слоя) — «скан»."""
    doc = pymupdf.open()
    for lines in pages:
        image = page_image(lines, photo)
        if rotate:
            image = image.rotate(rotate, expand=True, fillcolor=255)
        tmp = path.with_suffix(".tmp.png")
        image.save(tmp)
        page = doc.new_page(width=image.width * 72 / 200, height=image.height * 72 / 200)
        page.insert_image(page.rect, filename=str(tmp))
        tmp.unlink()
    doc.save(str(path), deflate=True)
    doc.close()


def make_docx(path: Path):
    """DOCX с абзацами и таблицей."""
    doc = Document()
    doc.add_heading("Протокол совещания", level=1)
    doc.add_paragraph("На совещании утверждён квартальный отчёт за третий квартал.")
    doc.add_paragraph("Следующий квартальный отчёт будет готов в январе.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Документ"
    table.cell(0, 1).text = "Статус"
    table.cell(1, 0).text = "Квартальный отчёт"
    table.cell(1, 1).text = "утверждён"
    doc.save(str(path))


def make_xlsx(path: Path):
    """XLSX с двумя листами."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Платежи"
    sheet.append(["Номер", "Описание", "Сумма"])
    sheet.append([4711, "Квартальный отчёт: оплата услуг", 12500])
    sheet.append([4712, "Канцелярские товары", 830])
    other = book.create_sheet("Заметки")
    other.append(["Напоминание: сдать квартальный отчёт до пятницы"])
    book.save(str(path))


def make_pptx(path: Path):
    """PPTX с заметками докладчика."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Итоги года"
    slide.placeholders[1].text = "Главное: квартальный отчёт принят"
    slide.notes_slide.notes_text_frame.text = "Не забыть упомянуть квартальный отчёт"
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Вопросы"
    deck.save(str(path))


def main():
    """Пересоздать демо-набор."""
    if DEMO.exists():
        for child in DEMO.iterdir():
            if child.name != "README.md":
                shutil.rmtree(child) if child.is_dir() else child.unlink()
    DEMO.mkdir(exist_ok=True)
    docs = DEMO / "документы"
    scans = DEMO / "сканы"
    arch = DEMO / "архивы"
    errors = DEMO / "для_проверки_ошибок"
    for folder in (docs, scans, arch, errors):
        folder.mkdir(exist_ok=True)
    text_pdf(
        docs / "01_текстовый_слой.pdf",
        [
            "Страница 1. Вводная часть без искомых слов.",
            "Страница 2. Здесь лежит квартальный отчёт компании.",
            "Страница 3. Заключение.",
        ],
    )
    text_pdf(
        docs / "03_повторы_на_странице.pdf",
        [
            "Квартальный отчёт: раздел 1.\nКвартальный отчёт: раздел 2.\n"
            "Итог: квартальный отчёт сдан."
        ],
    )
    make_docx(docs / "04_протокол.docx")
    make_xlsx(docs / "05_платежи.xlsx")
    make_pptx(docs / "06_презентация.pptx")
    (docs / "07_заметки.txt").write_text(
        "Первая строка.\nСдать Квартальный Отчёт до пятницы.\nКонец файла.\n", encoding="utf-8"
    )
    (docs / "08_журнал_работы.txt").write_text(
        "12:00 старт\n12:05 сформирован квартальный отчёт\n12:06 конец\n", encoding="utf-8"
    )
    (docs / "09_данные.csv").write_text(
        "id;название\n1;Квартальный отчёт\n2;Годовой план\n", encoding="cp1251"
    )
    (docs / "10_readme.md").write_text("# Папка\n\nВ ней нет искомых слов.\n", encoding="utf-8")
    (docs / "16_заглавные.txt").write_text(
        "КВАРТАЛЬНЫЙ ОТЧЁТ — ЗАГЛАВНЫМИ БУКВАМИ\n", encoding="utf-8"
    )
    (docs / "17_формат_4711-К.txt").write_text(
        "Договор 4711-К и договор 4712-К\n", encoding="utf-8"
    )
    scan_pdf(
        scans / "02_скан_книги.pdf",
        [
            ["Предисловие", "Эта страница не содержит нужных слов."],
            ["Глава вторая", "Квартальный отчёт утверждён.", "Конец главы."],
        ],
    )
    scan_pdf(
        scans / "18_фото_страницы.pdf",
        [["Фото документа", "Квартальный отчёт получен."]],
        photo=True,
    )
    scan_pdf(
        scans / "19_повёрнутый_скан.pdf",
        [["Повёрнутая страница", "Квартальный отчёт найден."]],
        rotate=90,
    )
    scan_pdf(scans / "15_пустая_страница.pdf", [[]])
    inner_scan = scans / "_tmp_scan.pdf"
    scan_pdf(inner_scan, [["Вложенный архив", "Квартальный отчёт внутри скана."]])
    nested = arch / "_nested.zip"
    with zipfile.ZipFile(nested, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(inner_scan, "скан_во_вложенном.pdf")
    with zipfile.ZipFile(arch / "11_архив_с_вложенным.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(docs / "04_протокол.docx", "папка/протокол.docx")
        zf.write(nested, "папка/вложенный.zip")
        zf.writestr("папка/заметка.txt", "Тут тоже есть квартальный отчёт.\n")
    nested.unlink()
    inner_scan.unlink()
    good = (arch / "11_архив_с_вложенным.zip").read_bytes()
    (errors / "12_битый_архив.zip").write_bytes(good[: len(good) // 2])
    with zipfile.ZipFile(errors / "14_выход_за_папку.zip", "w") as zf:
        zf.writestr("../evil.txt", "квартальный отчёт — запись должна быть отклонена\n")
        zf.writestr("нормальный.txt", "квартальный отчёт в безопасной записи\n")
    protected = pymupdf.open()
    protected.new_page().insert_text((60, 100), "секрет", fontname="dv", fontfile=FONT)
    protected.save(
        str(errors / "13_зашифрованный.pdf"),
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="user",
    )
    protected.close()
    total = sum(f.stat().st_size for f in DEMO.rglob("*") if f.is_file())
    count = sum(1 for f in DEMO.rglob("*") if f.is_file())
    print(f"Создано файлов: {count}, размер: {total / 1024:.0f} КБ")


if __name__ == "__main__":
    main()
